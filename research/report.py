"""The program report: figures, REPORT.md and a deck, into one dated folder.

    research/outputs/*.json -> figures -> REPORT.md + adaptive_therapy_report.pptx

Modelled on herc_llm's research/report.py. **Result numbers are read from run
artifacts, never typed in**; facts from before those artifacts existed (Phase
0-1, the Phase-2 PPO runs) cite the RESEARCH_LOG line that holds them. If an
artifact is missing, the section that needs it is replaced by a visible note,
never by a guess.

The dated folder is the unit of delivery: `report/YYYY-MM-DD/`. Re-running on
the same day overwrites that day's folder.

    JAX_PLATFORMS=cpu uv run --group report python research/report.py
"""

from __future__ import annotations

import json
import shutil
import subprocess
from datetime import date
from pathlib import Path

import matplotlib

matplotlib.use("Agg")
import matplotlib.pyplot as plt  # noqa: E402
import numpy as np  # noqa: E402
from PIL import Image  # noqa: E402
from pptx import Presentation  # noqa: E402
from pptx.chart.data import CategoryChartData  # noqa: E402
from pptx.dml.color import RGBColor  # noqa: E402
from pptx.enum.chart import XL_CHART_TYPE, XL_LEGEND_POSITION  # noqa: E402
from pptx.enum.shapes import MSO_SHAPE  # noqa: E402
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN  # noqa: E402
from pptx.util import Inches, Pt  # noqa: E402

REPO = Path(__file__).resolve().parents[1]
OUT = REPO / "research" / "outputs"
DECK = OUT / "deck"  # written by research/deck_figures.py

INK = RGBColor(0x1F, 0x3B, 0x4D)
MIST = RGBColor(0xEE, 0xF3, 0xF5)
CORAL = RGBColor(0xD1, 0x49, 0x5B)
TEAL = RGBColor(0x2A, 0x9D, 0x8F)
GREY = RGBColor(0x5F, 0x6B, 0x73)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
HEAD, BODY = "Cambria", "Calibri"
W, H, M = 13.333, 7.5, 0.6
REGIME_COLORS = {
    "none": RGBColor(0x9A, 0x9A, 0x9A),
    "MTD": CORAL,
    "AT50": RGBColor(0xE0, 0x9F, 0x3E),
    "thr0.8": RGBColor(0x3C, 0x6E, 0x71),
    "MLP": INK,
    "LSTM": RGBColor(0x8E, 0x7D, 0xBE),
    "ODE-RNN": TEAL,
}
ARCH_HEX = {"mlp": "#1F3B4D", "lstm": "#8E7DBE", "ode_rnn": "#2A9D8F"}


def load(name: str) -> dict | None:
    p = OUT / name
    return json.loads(p.read_text()) if p.exists() else None


# --- figures ------------------------------------------------------------------


def fig_validity(runs: dict, path: Path) -> None:
    """Closed-loop score of every trained cell, per registration and arm."""
    fig, ax = plt.subplots(figsize=(10, 4))
    xt = []
    for i, (reg, d) in enumerate(runs.items()):
        for j, arch in enumerate(("mlp", "lstm", "ode_rnn")):
            vals = [v for k, v in d.items() if k.startswith(arch + "_w")]
            if not vals:
                continue
            x = i * 4 + j
            jit = np.linspace(-0.25, 0.25, len(vals))
            ax.scatter(
                x + jit,
                vals,
                s=22,
                color=ARCH_HEX[arch],
                alpha=0.85,
                label=arch if i == len(runs) - 1 else None,
            )
            bad = [(xx, v) for xx, v in zip(x + jit, vals, strict=True) if v < 0.10]
            if bad:
                ax.scatter(
                    [b[0] for b in bad],
                    [b[1] for b in bad],
                    s=110,
                    facecolors="none",
                    edgecolors="#D1495B",
                    lw=1.5,
                )
        xt.append((i * 4 + 1, reg))
    ax.axhline(0.10, color="#D1495B", ls="--", lw=1)
    ax.text(
        ax.get_xlim()[1],
        0.105,
        "validity gate 0.10",
        ha="right",
        va="bottom",
        fontsize=9,
        color="#D1495B",
    )
    ax.set_xticks([x for x, _ in xt], [r for _, r in xt])
    ax.set_ylabel("closed-loop score (Tier 1)")
    ax.legend(frameon=False, fontsize=9, loc="upper right")
    ax.spines[["top", "right"]].set_visible(False)
    fig.tight_layout()
    fig.savefig(path, dpi=170)
    plt.close(fig)


def build_figures(fig_dir: Path, arts: dict) -> dict[str, Path]:
    fig_dir.mkdir(parents=True, exist_ok=True)
    figs = {}

    def put(key, src: Path, n: int):
        if src.exists():
            dst = fig_dir / f"{n:02d}_{key}.png"
            shutil.copyfile(src, dst)
            figs[key] = dst

    put("untreated", OUT / "phase1_untreated_trajectories.png", 1)
    for n, fam in enumerate(("lotka_volterra", "hahnfeldt", "kuznetsov"), start=2):
        put(f"regimes_{fam}", DECK / f"regimes_{fam}.png", n)
    put("mlp_rule", DECK / "mlp_rule.png", 5)
    put("w3_frontier", OUT / "w3_frontier.png", 6)
    put("w3b_frontier", OUT / "w3b_frontier.png", 7)
    runs = {
        name: d["observation_dependence" if name == "Phase 2c" else "closed_loop_tier1"]
        for name, d in (("Phase 2c", arts["p2c"]), ("W3", arts["w3"]), ("W3b", arts["w3b"]))
        if d
    }
    if runs:
        fig_validity(runs, fig_dir / "08_validity.png")
        figs["validity"] = fig_dir / "08_validity.png"
    return figs


# --- numbers read from artifacts ---------------------------------------------------


def numbers(arts: dict) -> dict:
    n: dict = {}
    if p := arts["p2c"]:
        for w, v in p["per_weight"].items():
            n[f"p2c_{w}"] = {
                "earlier": v["pairs_ode_earlier"],
                "later": v["pairs_ode_later"],
                "expo_ode": v["exposure_ode"],
                "expo_mlp": v["exposure_mlp"],
            }
    if d := arts["w1"]:
        n["w1"] = {k: (v["ode_earlier"], v["ode_later"]) for k, v in d["paired"].items()}
        cells = d["cells"]
        n["w1_obs_mlp"] = min(
            c["tier1"]["dose_explained_by_obs"] for c in cells if c["arch"] == "mlp"
        )
        ode = [c["tier1"]["dose_explained_by_obs"] for c in cells if c["arch"] == "ode_rnn"]
        n["w1_obs_ode"] = (min(ode), max(ode))
    for reg in ("w3", "w3b"):
        if d := arts[reg]:
            n[reg] = {
                "fails": {k: v for k, v in d["closed_loop_tier1"].items() if v < 0.10},
                "H_mem": d["H_mem"],
                "H_ct": d["H_ct"],
                "valid_arm": d["valid_arm"],
            }
    n["regimes"] = arts["regimes"]
    return n


# --- deck --------------------------------------------------------------------------


class Deck:
    def __init__(self):
        self.prs = Presentation()
        self.prs.slide_width, self.prs.slide_height = Inches(W), Inches(H)
        self.blank = self.prs.slide_layouts[6]

    def new(self, dark=False, notes=None):
        s = self.prs.slides.add_slide(self.blank)
        f = s.background.fill
        f.solid()
        f.fore_color.rgb = INK if dark else WHITE
        if notes:
            s.notes_slide.notes_text_frame.text = notes
        return s

    @staticmethod
    def text(
        slide,
        x,
        y,
        w,
        h,
        runs,
        size=16,
        color=INK,
        font=BODY,
        bold=False,
        align=PP_ALIGN.LEFT,
        italic=False,
        space_after=6,
    ):
        """`runs`: a string, or paragraphs; a paragraph is a string or
        a list of (text, {overrides}) tuples."""
        tb = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
        tf = tb.text_frame
        tf.word_wrap = True
        for m in ("margin_left", "margin_right", "margin_top", "margin_bottom"):
            setattr(tf, m, 0)
        for i, para in enumerate([runs] if isinstance(runs, str) else runs):
            p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
            p.alignment = align
            p.space_after = Pt(space_after)
            for t, o in [(para, {})] if isinstance(para, str) else para:
                r = p.add_run()
                r.text = t
                r.font.size = Pt(o.get("size", size))
                r.font.name = o.get("font", font)
                r.font.bold = o.get("bold", bold)
                r.font.italic = o.get("italic", italic)
                r.font.color.rgb = o.get("color", color)
        return tb

    def title(self, slide, t, sub=None, dark=False):
        self.text(
            slide,
            M,
            0.45,
            W - 2 * M,
            0.8,
            t,
            size=34,
            color=WHITE if dark else INK,
            font=HEAD,
            bold=True,
        )
        if sub:
            self.text(
                slide,
                M,
                1.2,
                W - 2 * M,
                0.5,
                sub,
                size=16,
                color=MIST if dark else GREY,
                italic=True,
            )

    @staticmethod
    def shape(slide, kind, x, y, w, h, fill):
        s = slide.shapes.add_shape(kind, Inches(x), Inches(y), Inches(w), Inches(h))
        s.fill.solid()
        s.fill.fore_color.rgb = fill
        s.line.fill.background()
        s.shadow.inherit = False
        return s

    def card(self, slide, x, y, w, h, fill=MIST):
        s = self.shape(slide, MSO_SHAPE.ROUNDED_RECTANGLE, x, y, w, h, fill)
        s.adjustments[0] = 0.06
        return s

    def badge(self, slide, x, y, label, fill=CORAL, d=0.55, size=16):
        s = self.shape(slide, MSO_SHAPE.OVAL, x, y, d, d, fill)
        tf = s.text_frame
        for m in ("margin_left", "margin_right", "margin_top", "margin_bottom"):
            setattr(tf, m, 0)
        tf.vertical_anchor = MSO_ANCHOR.MIDDLE
        p = tf.paragraphs[0]
        p.alignment = PP_ALIGN.CENTER
        r = p.add_run()
        r.text = label
        r.font.size, r.font.bold, r.font.name = Pt(size), True, BODY
        r.font.color.rgb = WHITE

    @staticmethod
    def image(slide, path, x, y, w, h):
        """Fit inside the box, preserving aspect, centred."""
        iw, ih = Image.open(path).size
        k = min(w / iw, h / ih)
        slide.shapes.add_picture(
            str(path),
            Inches(x + (w - iw * k) / 2),
            Inches(y + (h - ih * k) / 2),
            Inches(iw * k),
            Inches(ih * k),
        )

    def missing(self, slide, what):
        self.text(
            slide,
            M,
            3.0,
            W - 2 * M,
            1,
            f"Artifact missing: {what}. Section omitted.",
            size=18,
            color=CORAL,
            italic=True,
        )


def verdict_word(h: dict) -> str:
    return "PASS" if h["passes"] else ("fail" if h["valid"] else "INVALID")


def build_deck(path: Path, figs: dict, n: dict, stamp: str) -> None:
    d = Deck()
    T, card, badge, image = d.text, d.card, d.badge, d.image

    # 1 title
    s = d.new(
        dark=True,
        notes="Result numbers are read from run artifacts by research/report.py; "
        "earlier facts cite RESEARCH_LOG lines. Virtual patients only.",
    )
    T(
        s,
        M,
        2.0,
        11.5,
        1.6,
        "Does a continuous-time policy dose cancer therapy more robustly?",
        size=40,
        color=WHITE,
        font=HEAD,
        bold=True,
    )
    T(
        s,
        M,
        3.75,
        11.5,
        0.6,
        "Continuous-time RL for adaptive therapy — what we built, ran, and learned",
        size=20,
        color=MIST,
        italic=True,
    )
    T(
        s,
        M,
        6.3,
        11.5,
        0.4,
        f"Program report · {stamp} · virtual patients only, no clinical claims",
        size=14,
        color=MIST,
    )

    # 2 question / pipeline
    s = d.new(
        notes="The ablation axis is the policy architecture; everything else in the loop is fixed. "
        "The agent only sees total burden, never the sensitive/resistant split, so this is a POMDP."
    )
    d.title(
        s, "The question", "Is it the differential-equation structure that helps — or just memory?"
    )
    steps = [
        ("Virtual patient", "model + θ"),
        ("ODE solver", "diffrax, weekly"),
        ("Burden only", "noisy, σ = 0.05"),
        ("Policy", "MLP · LSTM · ODE-RNN"),
        ("Dose", "0 · ½ · 1"),
    ]
    bw, gap, y0 = 2.05, 0.42, 2.35
    for i, (a, b) in enumerate(steps):
        x = M + i * (bw + gap)
        card(s, x, y0, bw, 1.35, fill=INK if i == 3 else MIST)
        T(
            s,
            x + 0.15,
            y0 + 0.25,
            bw - 0.3,
            0.45,
            a,
            size=15,
            bold=True,
            color=WHITE if i == 3 else INK,
            align=PP_ALIGN.CENTER,
        )
        T(
            s,
            x + 0.15,
            y0 + 0.75,
            bw - 0.3,
            0.45,
            b,
            size=13,
            color=MIST if i == 3 else GREY,
            align=PP_ALIGN.CENTER,
        )
        if i < len(steps) - 1:
            d.shape(s, MSO_SHAPE.RIGHT_ARROW, x + bw + 0.07, y0 + 0.52, gap - 0.14, 0.3, GREY)
    for i, (a, b) in enumerate(
        [
            ("MLP", "memoryless — the null hypothesis"),
            ("LSTM", "memory, discrete steps"),
            ("ODE-RNN", "memory + continuous time between visits"),
        ]
    ):
        x = M + i * 4.1
        badge(s, x, 4.45, str(i + 1), fill=TEAL if i == 2 else INK)
        T(s, x + 0.75, 4.42, 3.2, 0.4, a, size=18, bold=True)
        T(s, x + 0.75, 4.85, 3.2, 0.7, b, size=14, color=GREY)
    T(
        s,
        M,
        6.25,
        W - 2 * M,
        0.6,
        [
            [
                ("Endpoints, always together: ", {"bold": True}),
                (
                    "time to progression (TTP: +20% over running nadir, 28-day confirmation) and cumulative drug "
                    "exposure. Paired by patient, ≥ 3 seeds.",
                    {},
                ),
            ]
        ],
        size=14,
    )

    # 3 equations
    s = d.new(
        notes="Three mechanistically distinct families. Treatment enters differently in each: it slows "
        "sensitive-cell growth (LV), destroys vasculature (Hahnfeldt), or kills tumour cells (Kuznetsov)."
    )
    d.title(s, "Three tumour models, three different biologies")
    models = [
        (
            "Lotka–Volterra",
            "competition",
            [
                "dS/dt = r_S S (1 − (S+R)/K)(1 − d_D u) − d_S S",
                "dR/dt = r_R R (1 − (S+R)/K) − d_R R",
            ],
            "Observed: S + R. Drug slows only sensitive cells; resistant cells are held back by competition "
            "— the mechanism adaptive therapy exploits.",
        ),
        (
            "Hahnfeldt",
            "angiogenesis",
            ["dV/dt = −a V ln(V/K)", "dK/dt = b V − d V^(2/3) K − e u K"],
            "Observed: V. The limit is the vascular capacity K, not competition. Drug attacks K.",
        ),
        (
            "Kuznetsov",
            "immune predation",
            ["dE/dt = s + pET/(g+T) − mET − δE", "dT/dt = aT(1 − bT) − nET − uT"],
            "Observed: T. Effector cells hunt the tumour; oscillation and dormancy are possible. Drug kills T.",
        ),
    ]
    cw = (W - 2 * M - 0.7) / 3
    for i, (name, kind, eqs, desc) in enumerate(models):
        x = M + i * (cw + 0.35)
        card(s, x, 1.5, cw, 5.3)
        T(s, x + 0.3, 1.75, cw - 0.6, 0.5, name, size=22, bold=True, font=HEAD)
        T(s, x + 0.3, 2.25, cw - 0.6, 0.4, kind, size=14, color=CORAL, italic=True)
        T(s, x + 0.3, 2.9, cw - 0.6, 1.3, eqs, size=15, font=HEAD, italic=True, space_after=10)
        T(s, x + 0.3, 4.45, cw - 0.6, 2.1, desc, size=14, color=GREY)

    # 4 untreated
    s = d.new(notes="Phase 0-1: RESEARCH_LOG L13, L48, L88.")
    d.title(
        s,
        "Untreated, they behave nothing alike",
        "Phase 0–1: families separable from burden alone (0.996, L48), but θ is not identifiable",
    )
    if "untreated" in figs:
        image(s, figs["untreated"], M, 1.95, W - 2 * M, 3.8)
    else:
        d.missing(s, "phase1_untreated_trajectories.png")
    T(
        s,
        M,
        6.1,
        W - 2 * M,
        0.9,
        "Untreated TTP at weekly visits: LV 28 d, Hahnfeldt 7 d, Kuznetsov 14 d "
        "(L88). Only LV has dynamic range, so the architecture comparison uses LV.",
        size=14,
        color=GREY,
    )

    # 5 regimes legend
    s = d.new(
        notes="Every regime runs through the same closed-loop rollout. The trained networks are "
        "LV-trained (w 0.05, seed 0) and only RUN on the other two models."
    )
    d.title(
        s,
        "Dosing regimes compared",
        "Same simulator loop, same weekly visit, only the decision rule changes",
    )
    regs = [
        ("none", "never treat"),
        ("MTD", "full dose every visit — standard of care"),
        ("AT50", "treat until burden < 50% of baseline, hold until back at baseline (Zhang trial)"),
        ("thr0.8", "treat whenever burden > 0.8 × baseline — the rule Gallagher 2024's RL found"),
        ("MLP", "trained, memoryless"),
        ("LSTM", "trained, discrete-time memory"),
        ("ODE-RNN", "trained, continuous-time memory"),
    ]
    for i, (a, b) in enumerate(regs):
        ci, row = divmod(i, 4)
        x, y = M + ci * 6.2, 1.95 + row * 1.2
        badge(s, x, y, "", fill=REGIME_COLORS[a], d=0.42)
        T(s, x + 0.65, y - 0.02, 5.3, 0.4, a, size=18, bold=True)
        T(s, x + 0.65, y + 0.38, 5.3, 0.7, b, size=14, color=GREY)

    # 6-8 per model
    per_model = [
        (
            "lotka_volterra",
            "LV: holding burden near 80% contains resistance",
            [
                "MLP and the 0.8 threshold rule cycle the drug and keep the tumour below baseline all year.",
                "AT50 lets it rebound to baseline before re-treating; MTD shrinks it fastest with the most drug.",
                "The LSTM switches to half dose and drifts; the ODE-RNN holds full dose, then half.",
            ],
        ),
        (
            "hahnfeldt",
            "Hahnfeldt: no regime changes the outcome",
            [
                "The tumour grows ~200× within weeks whatever the dose: every regime progresses on day 7.",
                "This is Phase 1's resolution limit (L88): treated TTP has no dynamic range here.",
                "W5 (mechanism shift) needs this fixed before it can be registered.",
            ],
        ),
        (
            "kuznetsov",
            "Kuznetsov: only full dose eradicates",
            [
                "LV-style rules treat once, see burden fall 100×, and stop while the immune balance is lost.",
                "MTD drives the tumour to extinction; the recurrent nets suppress it at very different drug.",
                "A threshold learned on one biology does not transfer to another.",
            ],
        ),
    ]
    for fam, t, bullets in per_model:
        s = d.new(
            notes="Illustration, not a registered result: the patient whose untreated TTP is closest to "
            "the median of 100 sampled. Networks trained on LV only. research/deck_figures.py."
        )
        d.title(s, t)
        if f"regimes_{fam}" in figs:
            image(s, figs[f"regimes_{fam}"], M, 1.35, 7.7, 5.75)
        else:
            d.missing(s, f"regimes_{fam}.png")
        for i, b in enumerate(bullets):
            y = 1.75 + i * 1.7
            badge(s, 8.75, y, str(i + 1), fill=INK, d=0.45, size=14)
            T(s, 9.4, y - 0.02, 3.35, 1.5, b, size=14)

    # 9 native chart
    s = d.new(
        notes="100 virtual patients per model, 52 weekly visits; restricted-mean TTP = mean of TTP "
        "capped at 357 d. Illustration, not registered."
    )
    d.title(
        s,
        "Across 100 patients: TTP bought, and at what dose",
        "Restricted-mean TTP (days, capped at 357) · illustration, not a registered comparison",
    )
    reg = n["regimes"]
    if reg:
        fams = ("lotka_volterra", "hahnfeldt", "kuznetsov")
        cd = CategoryChartData()
        cd.categories = ["Lotka–Volterra", "Hahnfeldt", "Kuznetsov"]
        for name in REGIME_COLORS:
            cd.add_series(name, [round(reg[f][name]["ttp_rmst"]) for f in fams])
        ch = s.shapes.add_chart(
            XL_CHART_TYPE.COLUMN_CLUSTERED, Inches(M), Inches(1.9), Inches(8.2), Inches(5.1), cd
        ).chart
        ch.has_legend = True
        ch.legend.position = XL_LEGEND_POSITION.TOP
        ch.legend.include_in_layout = False
        ch.legend.font.size = Pt(11)
        for ser, col in zip(ch.series, REGIME_COLORS.values(), strict=True):
            ser.format.fill.solid()
            ser.format.fill.fore_color.rgb = col
        ch.value_axis.maximum_scale = 360
        ch.value_axis.major_gridlines.format.line.color.rgb = RGBColor(0xDD, 0xDD, 0xDD)
        ch.value_axis.tick_labels.font.size = Pt(11)
        ch.category_axis.tick_labels.font.size = Pt(12)
        lv = reg["lotka_volterra"]
        card(s, 9.2, 1.9, 3.55, 5.1)
        saving = 1 - lv["MLP"]["exposure_median"] / lv["MTD"]["exposure_median"]
        T(
            s,
            9.45,
            2.1,
            3.1,
            4.8,
            [
                [("Lotka–Volterra", {"bold": True})],
                f"MLP {lv['MLP']['ttp_rmst']:.0f} d at exposure {lv['MLP']['exposure_median']:.0f}",
                f"thr0.8 {lv['thr0.8']['ttp_rmst']:.0f} d at {lv['thr0.8']['exposure_median']:.0f}",
                f"MTD {lv['MTD']['ttp_rmst']:.0f} d at {lv['MTD']['exposure_median']:.0f}",
                [
                    (
                        f"→ the trained MLP ≈ a one-line rule, with {saving:.0%} less drug than MTD.",
                        {"color": CORAL},
                    )
                ],
                [("Hahnfeldt", {"bold": True})],
                f"{reg['hahnfeldt']['MTD']['ttp_rmst']:.0f} d for every regime: no signal at weekly visits.",
                [("Kuznetsov", {"bold": True})],
                f"only MTD reaches the horizon ({reg['kuznetsov']['MTD']['ttp_rmst']:.0f} d).",
            ],
            size=13,
            space_after=5,
        )
    else:
        d.missing(s, "deck/summary.json (run research/deck_figures.py)")

    # 10 timeline
    s = d.new(
        notes="Every comparison pre-registered in PREREG.md and logged in RESEARCH_LOG.md whatever the "
        "outcome. Void and invalid runs stay in the record."
    )
    d.title(
        s,
        "What was run, in order",
        "Every comparison pre-registered; void and invalid runs kept in the record",
    )
    w3b_state = "INVALID" if n.get("w3b") and not all(n["w3b"]["valid_arm"].values()) else "DONE"
    tl = [
        ("Phase 0", "families separable\nfrom burden", "PASS", TEAL),
        ("Phase 1", "simulator +\nTTP rule", "DONE", TEAL),
        ("Phase 2", "shifted cohort\nmis-built", "VOID", GREY),
        ("Phase 2 / 2b", "PPO: ODE-RNN\nopen-loop", "INVALID", GREY),
        ("Phase 2c", "PQN: first valid\nrun", "FAIL", CORAL),
        ("W1", "diagnosis on\nsaved policies", "DONE", TEAL),
        ("W3", "+ LSTM arm,\nfrontiers", "INVALID", GREY),
        ("W3b", "fresh seeds,\nRMST", w3b_state, GREY),
    ]
    d.shape(
        s, MSO_SHAPE.RECTANGLE, M + 0.3, 3.62, W - 2 * M - 0.6, 0.04, RGBColor(0xC9, 0xD3, 0xD8)
    )
    step = (W - 2 * M) / len(tl)
    for i, (a, b, st, col) in enumerate(tl):
        cx = M + step * i + step / 2
        badge(s, cx - 0.28, 3.36, "", fill=col, d=0.56)
        T(
            s,
            cx - step / 2 + 0.05,
            2.35,
            step - 0.1,
            0.45,
            a,
            size=15,
            bold=True,
            align=PP_ALIGN.CENTER,
        )
        T(
            s,
            cx - step / 2 + 0.05,
            2.8,
            step - 0.1,
            0.4,
            st,
            size=12,
            bold=True,
            color=col,
            align=PP_ALIGN.CENTER,
        )
        T(
            s,
            cx - step / 2 + 0.05,
            4.2,
            step - 0.1,
            0.9,
            b,
            size=12,
            color=GREY,
            align=PP_ALIGN.CENTER,
        )
    T(
        s,
        M,
        5.6,
        W - 2 * M,
        1.2,
        [
            [
                ("Why so many invalid runs? ", {"bold": True}),
                (
                    "Each caught a way this evaluation can lie: a policy that ignores its input, a median that rewards "
                    "letting patients progress, a cohort built from the unshifted parameter, a metric stuck at its "
                    "ceiling. Each is now a gate in every run.",
                    {},
                ),
            ]
        ],
        size=14,
    )

    # 11 phase 2c
    s = d.new(
        notes="L424. Paired by patient, pooled over 3 seeds x 200 patients. Caveat: the MLP doses more, "
        "and on LV more drug alone lengthens TTP, so 'worse at equal drug' is not established here."
    )
    d.title(
        s,
        "Phase 2c: the first valid run — the ODE-RNN loses",
        "LV, Tier-1 shift (×1.5 growth rates, ×0.6 capacity), PQN for both arms, 3 seeds × 200 patients",
    )
    if "p2c_0.05" in n:
        a, b = n["p2c_0.05"], n["p2c_0.2"]
        stats = [
            (
                f"{a['earlier']} vs {a['later']}",
                "patients progressing earlier vs later on the ODE-RNN",
                "exposure weight 0.05",
            ),
            (
                f"{b['earlier']} vs {b['later']}",
                "patients progressing earlier vs later on the ODE-RNN",
                "exposure weight 0.20",
            ),
            (
                f"{a['expo_ode'] / a['expo_mlp'] - 1:+.0%}",
                f"drug given by the ODE-RNN ({a['expo_ode']:.0f} vs {a['expo_mlp']:.0f})",
                "so TTP is not exposure-matched",
            ),
        ]
        for i, (big, lab, sub) in enumerate(stats):
            x = M + i * 4.1
            card(s, x, 2.0, 3.75, 3.2)
            T(
                s,
                x + 0.3,
                2.25,
                3.15,
                1.1,
                big,
                size=48,
                bold=True,
                font=HEAD,
                color=CORAL if i < 2 else INK,
            )
            T(s, x + 0.3, 3.45, 3.15, 1.0, lab, size=15)
            T(s, x + 0.3, 4.5, 3.15, 0.5, sub, size=13, color=GREY, italic=True)
    else:
        d.missing(s, "phase2c_gate.json")
    T(
        s,
        M,
        5.65,
        W - 2 * M,
        1.2,
        [
            [
                ("Kill criterion fired: ", {"bold": True, "color": CORAL}),
                (
                    "no continuous-time edge on the clean single-model case, with either trainer (PPO or PQN). "
                    "The multi-model curriculum was not built.",
                    {},
                ),
            ]
        ],
        size=16,
    )

    # 12 validity
    s = d.new(
        notes="Closed-loop score = fraction of patient-visits whose greedy dose changes when another "
        "patient's observation history is swapped in. 0 = one fixed schedule for everyone."
    )
    d.title(
        s,
        "The recurrent arms sometimes stop listening",
        "Closed-loop score of every trained cell; circled cells fail the 0.10 gate",
    )
    if "validity" in figs:
        image(s, figs["validity"], M, 1.9, 8.4, 4.6)
    fails = {k: v for r in ("w3", "w3b") if n.get(r) for k, v in n[r]["fails"].items()}
    card(s, 9.3, 1.9, 3.45, 4.6)
    T(
        s,
        9.55,
        2.1,
        3.0,
        4.3,
        [[("Below the gate", {"bold": True, "size": 17})]]
        + [f"{k.replace('_w', ' w').replace('_s', ' s')}: {v:.3f}" for k, v in fails.items()]
        + [
            [
                (
                    "The MLP never fails; one recurrent cell in ~14 does — enough to invalidate every "
                    "registration with an LSTM arm so far.",
                    {"color": CORAL},
                )
            ]
        ],
        size=13,
        space_after=7,
    )

    # 13 W1
    s = d.new(notes="L461. Tier 0 = the same test patients without the shift.")
    d.title(
        s,
        "W1: it is a training gap, not a robustness gap",
        "The same policies on the patients before the shift (Tier 0)",
    )
    if "w1" in n:
        e0, l0 = n["w1"]["w0.05_tier0"]
        e1, l1 = n["w1"]["w0.05_tier1"]
        card(s, M, 1.95, 4.6, 4.9)
        T(
            s,
            M + 0.3,
            2.2,
            4.0,
            4.5,
            [
                [("Tier 0 (no shift)", {"bold": True, "size": 18})],
                [(f"{e0} vs {l0}", {"bold": True, "size": 36, "font": HEAD, "color": CORAL})],
                "ODE-RNN earlier vs later (w 0.05)",
                [("Tier 1 (shifted)", {"bold": True, "size": 18})],
                [(f"{e1} vs {l1}", {"bold": True, "size": 36, "font": HEAD, "color": CORAL})],
                "the shift moves both arms, not the gap",
            ],
            size=14,
        )
        lo, hi = n["w1_obs_ode"]
        T(
            s,
            5.6,
            5.75,
            7.1,
            1.1,
            [
                [
                    ("The MLP learned a switch: ", {"bold": True}),
                    (
                        f"no drug below ~0.8 × baseline, full dose above, in all 6 seeds (≥ {n['w1_obs_mlp']:.0%} of doses "
                        f"explained by the current reading). The ODE-RNN doses partly from memory ({lo:.0%}–{hi:.0%}).",
                        {},
                    ),
                ]
            ],
            size=14,
        )
    if "mlp_rule" in figs:
        image(s, figs["mlp_rule"], 5.6, 1.95, 7.1, 3.65)

    # 14 literature
    s = d.new(notes="Both read from the PDFs with page locators (literature/papers/).")
    d.title(
        s,
        "What the two closest papers actually show",
        "Read in full from the PDFs, with page locators",
    )
    lit = [
        (
            "Gallagher et al. 2024, Cancer Research",
            TEAL,
            [
                "Same setting: LV, burden-only, TTP = +20%.",
                "Their RL policy is a switch at N* = 0.83 × n₀ (Fig. 4B, p.1935) — our MLP found ~0.80.",
                "Evaluating less often than trained lowers TTP (Fig. 3C, p.1934): predicts W4's direction.",
                "Means over 100 evaluations; exposure not matched.",
            ],
        ),
        (
            "Kiram et al. 2026, arXiv 2605.02552",
            CORAL,
            [
                "Recurrent TD3 380 ± 12 vs TD3 −50 ± 252 (Table I, p.5).",
                "But: one fixed patient, no shift, means over 10 seeds.",
                "The hidden variable is a toxicity state in the reward — not our S/R split.",
                "Justifies an LSTM arm; not evidence that memory helps here.",
            ],
        ),
    ]
    for i, (t, col, bl) in enumerate(lit):
        x = M + i * 6.2
        card(s, x, 1.95, 5.85, 4.9)
        badge(s, x + 0.3, 2.2, "", fill=col, d=0.4)
        T(s, x + 0.9, 2.2, 4.7, 0.5, t, size=17, bold=True, font=HEAD)
        T(
            s,
            x + 0.3,
            3.0,
            5.25,
            3.7,
            [[("•  ", {"color": col, "bold": True}), (b, {})] for b in bl],
            size=14,
            space_after=12,
        )

    # 15-16 W3 / W3b
    for reg_key, t, sub, fig_key, ylab in (
        (
            "w3",
            "W3: adding an LSTM arm",
            "MLP vs LSTM isolates memory; LSTM vs ODE-RNN isolates continuous time",
            "w3_frontier",
            "capped median",
        ),
        (
            "w3b",
            "W3b: fresh seeds, restricted-mean TTP",
            "Same design; the ceiling removed",
            "w3b_frontier",
            "RMST",
        ),
    ):
        s = d.new(
            notes=f"{reg_key.upper()}: frontier = TTP ({ylab}) vs exposure across exposure weights; gaps are "
            "mean differences over the shared exposure range with a paired patient bootstrap CI."
        )
        d.title(s, t, sub)
        r = n.get(reg_key)
        if not r:
            d.missing(s, f"{reg_key}_frontier.json")
            continue
        if fig_key in figs:
            image(s, figs[fig_key], M, 1.9, 8.3, 3.4)
        rows = []
        for h, lab in (("H_mem", "Memory: LSTM − MLP"), ("H_ct", "Continuous time: ODE − LSTM")):
            t1 = r[h]["tier1"]
            rows += [
                [(lab, {"bold": True, "size": 15})],
                f"Tier 1 {t1['gap_days']:+.0f} d [{t1['ci95'][0]:+.0f}, {t1['ci95'][1]:+.0f}]",
                f"overlap {t1['overlap_fraction']:.0%} · {verdict_word(r[h])}",
            ]
        card(s, 9.2, 1.9, 3.55, 5.0)
        T(s, 9.45, 2.1, 3.1, 4.7, rows, size=14, space_after=7)
        bad = ", ".join(a for a, v in r["valid_arm"].items() if not v) or "none"
        T(
            s,
            M,
            5.55,
            8.3,
            1.4,
            [
                [
                    ("No verdict: ", {"bold": True, "color": CORAL}),
                    (f"arms with an open-loop cell: {bad}. Gaps are descriptive only.", {}),
                ]
            ]
            if bad != "none"
            else "All arms valid.",
            size=14,
        )

    # 17 pitfalls
    s = d.new(notes="ROADMAP traps 4-6 plus the W3 ceiling; each now has an automatic check.")
    d.title(s, "Four ways this evaluation lied — each now a gate")
    pit = [
        (
            "An open-loop schedule",
            "The ODE-RNN gave all 200 patients one identical dosing history and still "
            "reported a respectable TTP (L312).",
            "closed-loop validity ≥ 0.10 per cell",
        ),
        (
            "The progressor-only median",
            "Averaging only patients who progress rewards letting more of them "
            "progress; it inverted the paired result 3 times.",
            "min(TTP, horizon) over all patients",
        ),
        (
            "A mis-built shifted cohort",
            "Scaling K after building y0 put half the test cohort above capacity; "
            "NaN blow-ups scored as 'never progressed' (L242).",
            "rebuild y0 from shifted θ; refuse NaN",
        ),
        (
            "A metric at its ceiling",
            "Once > 50% never progress, a capped median is flat at the horizon and "
            "ranks nothing (L496).",
            "restricted-mean TTP (W3b)",
        ),
    ]
    for i, (a, b, fix) in enumerate(pit):
        x, y = M + (i % 2) * 6.2, 1.55 + (i // 2) * 2.75
        card(s, x, y, 5.85, 2.5)
        badge(s, x + 0.3, y + 0.3, str(i + 1), fill=CORAL, d=0.5)
        T(s, x + 1.0, y + 0.32, 4.6, 0.5, a, size=18, bold=True, font=HEAD)
        T(s, x + 0.3, y + 0.95, 5.25, 0.95, b, size=13, color=GREY)
        T(
            s,
            x + 0.3,
            y + 1.9,
            5.25,
            0.45,
            [[("Fix: ", {"bold": True, "color": TEAL}), (fix, {})]],
            size=13,
        )

    # 18 next
    s = d.new(
        dark=True, notes="Target: an ICLR 2027 workshop; a clean negative is still the paper."
    )
    d.title(
        s, "Next", "Target: ICLR 2027 workshop — a clean negative is still the paper", dark=True
    )
    nxt = [
        ("Rule", "decide per-cell vs per-arm validity before any new registration"),
        (
            "W4",
            "irregular visits (20% / 50% dropped) — where continuous time has a mechanical reason to help; "
            "LSTM + Δt as the strong discrete baseline",
        ),
        ("W5", "mechanism shift — only if Hahnfeldt's day-7 resolution limit is fixed"),
        (
            "W6",
            "write-up: validity across trainers, frontiers, the learned threshold rule, every run in one table",
        ),
    ]
    for i, (a, b) in enumerate(nxt):
        y = 2.1 + i * 1.15
        badge(s, M, y, a, fill=TEAL if i == 1 else CORAL, d=0.75, size=13)
        T(s, M + 1.05, y + 0.12, 10.8, 0.9, b, size=17, color=WHITE)

    d.prs.save(path)


# --- markdown -------------------------------------------------------------------


def build_markdown(path: Path, figs: dict, n: dict, stamp: str, commit: str) -> None:
    rel = {k: f"figures/{v.name}" for k, v in figs.items()}
    L = [
        "# Continuous-time RL for adaptive therapy — program report",
        "",
        f"**{stamp}** · commit `{commit}` · generated by `research/report.py`; result numbers are read "
        "from `research/outputs/` artifacts, earlier facts cite RESEARCH_LOG lines.",
        "",
        "---",
        "",
        "## 1. Summary",
        "",
        "**No continuous-time edge has been found, and no registered run has yet produced a valid "
        "memory-vs-continuous-time verdict.** On the one valid comparison (Phase 2c, L424) the ODE-RNN "
        "progresses earlier than a memoryless MLP, and W1 (L461) shows it loses in-distribution too — a "
        "training gap, not a robustness gap. Both LSTM registrations (W3 L496, W3b L526) are invalid: the "
        "recurrent arms go open-loop in about 1 in 14 PQN cells, the MLP never. Descriptively the MLP leads "
        "everywhere, and what it learned is a burden-threshold switch — the rule Gallagher 2024's RL found.",
        "",
    ]
    if "untreated" in rel:
        L += [
            "## 2. Three tumour models",
            "",
            f"![Untreated trajectories]({rel['untreated']})",
            "",
            "Lotka–Volterra (competition), Hahnfeldt (angiogenesis), Kuznetsov (immune predation). Families "
            "are separable from burden alone (0.996, L48) but θ is not identifiable. Only LV has untreated "
            "TTP dynamic range at weekly visits (L88).",
            "",
        ]
    L += [
        "## 3. Dosing regimes on each model (illustration)",
        "",
        "none · MTD · AT50 (Zhang trial rule) · thr0.8 (Gallagher's learned switch) · the LV-trained MLP, "
        "LSTM and ODE-RNN, all through the same closed-loop rollout.",
        "",
    ]
    for fam in ("lotka_volterra", "hahnfeldt", "kuznetsov"):
        if f"regimes_{fam}" in rel:
            L += [f"![{fam}]({rel[f'regimes_{fam}']})", ""]
    if n["regimes"]:
        L += [
            "| model | " + " | ".join(REGIME_COLORS) + " |",
            "|---" * (len(REGIME_COLORS) + 1) + "|",
        ]
        for fam, rows in n["regimes"].items():
            L.append(
                f"| {fam} | "
                + " | ".join(
                    f"{rows[k]['ttp_rmst']:.0f} d / {rows[k]['exposure_median']:.0f}"
                    for k in REGIME_COLORS
                )
                + " |"
            )
        L += [
            "",
            "*Restricted-mean TTP (capped 357 d) / median cumulative exposure, 100 patients per model.*",
            "",
        ]
    if "p2c_0.05" in n:
        a, b = n["p2c_0.05"], n["p2c_0.2"]
        L += [
            "## 4. Phase 2c — the first valid comparison (L424)",
            "",
            f"ODE-RNN earlier / later per matched patient: **{a['earlier']} / {a['later']}** (w 0.05), "
            f"**{b['earlier']} / {b['later']}** (w 0.20). Exposure ODE {a['expo_ode']:.0f} vs MLP "
            f"{a['expo_mlp']:.0f} — not matched. Kill criterion fired.",
            "",
        ]
    if "w1" in n:
        e0, l0 = n["w1"]["w0.05_tier0"]
        L += [
            "## 5. W1 — a training gap, not a robustness gap (L461)",
            "",
            f"On unshifted patients the ODE-RNN is still earlier in **{e0} / {e0 + l0}** untied pairs. The MLP "
            f"is a switch at ~0.8 × baseline (≥ {n['w1_obs_mlp']:.0%} of doses from the current reading).",
            "",
        ]
        if "mlp_rule" in rel:
            L += [f"![MLP rule]({rel['mlp_rule']})", ""]
    for reg_key, lab, line in (("w3", "W3", "L496"), ("w3b", "W3b", "L526")):
        r = n.get(reg_key)
        if not r:
            continue
        L += [
            f"## {'6' if reg_key == 'w3' else '7'}. {lab} — memory vs continuous time ({line})",
            "",
        ]
        if f"{reg_key}_frontier" in rel:
            L += [f"![{lab} frontier]({rel[f'{reg_key}_frontier']})", ""]
        L += [
            "| comparison | Tier 0 gap | Tier 1 gap | overlap T1 | status |",
            "|---|---|---|---|---|",
        ]
        for h, lab2 in (("H_mem", "LSTM − MLP"), ("H_ct", "ODE − LSTM")):
            t0, t1 = r[h]["tier0"], r[h]["tier1"]
            L.append(
                f"| {lab2} | {t0['gap_days']:+.0f} [{t0['ci95'][0]:+.0f}, {t0['ci95'][1]:+.0f}] | "
                f"{t1['gap_days']:+.0f} [{t1['ci95'][0]:+.0f}, {t1['ci95'][1]:+.0f}] | "
                f"{t1['overlap_fraction']:.0%} | {verdict_word(r[h])} |"
            )
        L += [
            "",
            "Open-loop cells: "
            + (", ".join(f"`{k}` {v:.3f}" for k, v in r["fails"].items()) or "none"),
            "",
        ]
    if "validity" in rel:
        L += [
            "## 8. Validity across registrations",
            "",
            f"![validity]({rel['validity']})",
            "",
            "Closed-loop score per trained cell; below 0.10 a policy is effectively one fixed schedule.",
            "",
        ]
    L += [
        "## 9. Literature read to 1°",
        "",
        "- **Gallagher et al. 2024** (Cancer Res, doi 10.1158/0008-5472.CAN-23-2040): RL on burden-only LV "
        "learns a switch at N* = 0.83 n₀ (Fig. 4B, p.1935); less frequent evaluation than training lowers TTP "
        "(Fig. 3C, p.1934).",
        "- **Kiram et al. 2026** (arXiv 2605.02552): recurrent TD3 wins when a toxicity state is hidden, on "
        "one fixed patient, means over 10 seeds (Table I, p.5).",
        "",
        "## 10. Next",
        "",
        "1. Decide per-cell vs per-arm validity before any new registration.",
        "2. W4 — irregular observation, with an LSTM + Δt baseline.",
        "3. W5 only if Hahnfeldt's day-7 resolution limit is fixed; W6 write-up.",
        "",
    ]
    path.write_text("\n".join(L))


def main() -> int:
    stamp = date.today().isoformat()
    folder = REPO / "report" / stamp
    if folder.exists():
        shutil.rmtree(folder)
    arts = {
        "p2c": load("phase2c_gate.json"),
        "w1": load("w1_diagnose.json"),
        "w3": load("w3_frontier.json"),
        "w3b": load("w3b_frontier.json"),
        "regimes": load("deck/summary.json"),
    }
    commit = subprocess.run(
        ["git", "rev-parse", "--short", "HEAD"], cwd=REPO, capture_output=True, text=True
    ).stdout.strip()
    figs = build_figures(folder / "figures", arts)
    n = numbers(arts)
    build_deck(folder / "adaptive_therapy_report.pptx", figs, n, stamp)
    build_markdown(folder / "REPORT.md", figs, n, stamp, commit)
    missing = [k for k, v in arts.items() if v is None]
    print(f"-> {folder}  ({len(figs)} figures; missing artifacts: {missing or 'none'})")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
